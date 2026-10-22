"""Validate final timeline package, editable objects and rendered text bounds."""
from pathlib import Path
from zipfile import ZipFile
import json, re, unicodedata, posixpath, hashlib, argparse
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import fitz

R=Path(__file__).resolve().parent
ppt=R/'LLM_Agent_Timeline.pptx';pdf=R/'LLM_Agent_Timeline.pdf'
p=Presentation(ppt);doc=fitz.open(pdf);issues=[]
norm=lambda s:''.join(unicodedata.normalize('NFKC',s).split())
page=doc[0];chars=[]
for b in page.get_text('rawdict')['blocks']:
 for li in b.get('lines',[]):
  for sp in li['spans']:
   for c in sp['chars']:
    for cc in norm(c['c']):chars.append((cc,c['bbox']))
alltext=''.join(c for c,_ in chars)
text_bounds=[]
for sh in p.slides[0].shapes:
 if sh.left<0 or sh.top<0 or sh.left+sh.width>p.slide_width+2 or sh.top+sh.height>p.slide_height+2:issues.append({'type':'shape_outside_page','name':sh.name})
 if not sh.has_text_frame or not sh.text.strip():continue
 expected=norm(sh.text)
 hits=[m.start() for m in re.finditer(re.escape(expected),alltext)]
 if not hits:issues.append({'type':'text_missing_in_pdf','text':sh.text});continue
 target=fitz.Rect(sh.left/12700,sh.top/12700,(sh.left+sh.width)/12700,(sh.top+sh.height)/12700)
 candidates=[]
 for st in hits:
  bb=fitz.Rect(chars[st][1])
  for _,r in chars[st+1:st+len(expected)]:bb|=fitz.Rect(r)
  candidates.append((abs(bb.x0-target.x0)+abs(bb.y0-target.y0),bb))
 bb=min(candidates,key=lambda a:a[0])[1]
 text_bounds.append({'text':sh.text,'bbox':list(bb),'shape_box':list(target)})
 if bb.x0<target.x0-2 or bb.x1>target.x1+2 or bb.y1>target.y1+3:issues.append({'type':'text_outside_own_box','text':sh.text,'bbox':list(bb),'shape_box':list(target)})
 if bb.x0<0 or bb.y0<0 or bb.x1>page.rect.width or bb.y1>page.rect.height:issues.append({'type':'text_outside_page','text':sh.text})
# Compare rendered text rectangles from different objects. Intentional dates/lines do not count as text overlaps.
for i,a in enumerate(text_bounds):
 for b in text_bounds[i+1:]:
  overlap=fitz.Rect(a['bbox'])&fitz.Rect(b['bbox'])
  if not overlap.is_empty and overlap.width>2 and overlap.height>2:issues.append({'type':'rendered_text_overlap','a':a['text'],'b':b['text'],'intersection':list(overlap)})
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
counts={'groups':0,'chart_parts':0,'embeddings':0,'ole':0,'optional_extensions':0}
with ZipFile(ppt) as z:
 assert z.testzip() is None
 names=z.namelist();assert len(names)==len(set(names));parts=set(names)
 parsed={n:etree.fromstring(z.read(n)) for n in names if n.endswith(('.xml','.rels'))}
 for n,x in parsed.items():
  if n.endswith('.rels'):
   source='' if n=='_rels/.rels' else posixpath.join(posixpath.dirname(posixpath.dirname(n)),posixpath.basename(n)[:-5])
   ids=[r.get('Id') for r in x]
   if len(ids)!=len(set(ids)):issues.append({'type':'duplicate_relationship_id','part':n})
   for r in x:
    if r.get('TargetMode')=='External':continue
    target=r.get('Target');target=posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(posixpath.dirname(source),target))
    if target not in parts:issues.append({'type':'missing_relationship_target','part':n,'target':target})
  else:
   relfile=posixpath.join(posixpath.dirname(n),'_rels',posixpath.basename(n)+'.rels')
   relids={r.get('Id') for r in parsed.get(relfile,[])}
   for el in x.iter():
    for k,v in el.attrib.items():
     if k.startswith('{'+NS['r']+'}') and v not in relids:issues.append({'type':'unresolved_relationship_reference','part':n,'id':v})
   ids=x.xpath('//p:cNvPr/@id',namespaces=NS)
   if len(ids)!=len(set(ids)):issues.append({'type':'duplicate_shape_id','part':n})
   counts['groups']+=len(x.findall('.//p:grpSp',NS));counts['ole']+=len(x.findall('.//p:oleObj',NS));counts['optional_extensions']+=len(x.xpath('//*[local-name()="extLst"]'))
 counts['chart_parts']=len([n for n in names if n.startswith('ppt/charts/')]);counts['embeddings']=len([n for n in names if n.startswith('ppt/embeddings/')])
works=json.loads((R/'timeline_data.json').read_text());layout=json.loads((R/'layout_manifest.json').read_text())
checks={
 'one_slide':len(p.slides)==1,'one_pdf_page':len(doc)==1,'16_9':abs(p.slide_width/p.slide_height-16/9)<1e-6,
 'nine_selected_works':len(works)==9,'all_selected_from_original':all(w['in_original'] for w in works),
 'chronological_months':all(a['date']<=b['date'] for a,b in zip(works,works[1:])),
 'month_labels_match_nodes':all(a['name']==b['name'] and a['date']==b['month'] for a,b in zip(works,layout)),
 'basic_objects_only':all(sh.shape_type in {MSO_SHAPE_TYPE.AUTO_SHAPE,MSO_SHAPE_TYPE.TEXT_BOX,MSO_SHAPE_TYPE.LINE,MSO_SHAPE_TYPE.PICTURE} for sh in p.slides[0].shapes),
 'no_complex_parts':not any(counts.values()),'all_notes_present':p.slides[0].has_notes_slide,
 'icons_match_manifest':len([sh for sh in p.slides[0].shapes if sh.shape_type==MSO_SHAPE_TYPE.PICTURE])==sum(bool(w['logo']) for w in works),
}
parser=argparse.ArgumentParser()
parser.add_argument('--roundtrip',type=Path,default=R/'roundtrip/LLM_Agent_Timeline.pptx')
args=parser.parse_args()
rt=args.roundtrip
if rt.exists():
 r=Presentation(rt)
 visible=lambda slide:'|'.join(norm(sh.text) for sh in slide.shapes if sh.has_text_frame)
 checks['libreoffice_roundtrip_one_slide']=len(r.slides)==1
 checks['libreoffice_roundtrip_text_equal']=visible(p.slides[0])==visible(r.slides[0])
 checks['libreoffice_roundtrip_notes_equal']=norm(p.slides[0].notes_slide.notes_text_frame.text)==norm(r.slides[0].notes_slide.notes_text_frame.text)
report={'passed':all(checks.values()) and not issues,'checks':checks,'complex_part_counts':counts,'issues':issues,'pptx_sha256':hashlib.sha256(ppt.read_bytes()).hexdigest(),'scope':'python-pptx reopen, ZIP/XML/relationships, PDF text bounds and overlaps; LibreOffice roundtrip if present. No Microsoft PowerPoint client test.'}
(R/'validation_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
if not report['passed']:raise SystemExit(1)
