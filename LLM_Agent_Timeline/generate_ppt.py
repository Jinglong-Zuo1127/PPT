#!/usr/bin/env python3
"""Generate the single-slide timeline using standard python-pptx objects.
Run: python generate_ppt.py
Optional preview: python generate_ppt.py --render
Dependencies: python-pptx, Pillow. Preview also needs LibreOffice and PyMuPDF.
"""
from pathlib import Path
import argparse, json, subprocess, os, tempfile
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'assets'
BLUE, RED, INK, GRAY = '2E5696', 'C62229', '222222', '666666'
FONT = '微软雅黑'
W, H = 13.333333333333334, 7.5

# Month-level dates avoid conflating preprints with later conference publication.
WORKS = [
 {'name':'ReAct','date':'2022.10','kind':'论文v1','label':'推理 × 行动','body':'交替组织推理与行动\n利用观察更新决策','logo':None,'name_text':'ReAct','paper':'https://arxiv.org/abs/2210.03629','code':'https://github.com/ysymyth/ReAct','venue':'ICLR 2023','in_original':True},
 {'name':'AutoGPT','date':'2023.03','kind':'官方演示','label':'自主连续执行','body':'任务分解与工具调用\n支持连续自主执行','logo':'autogpt_official_logo.png','name_text':'AutoGPT','paper':None,'code':'https://github.com/Significant-Gravitas/AutoGPT','venue':'开源项目；官方演示2023-03-30','in_original':True},
 {'name':'CAMEL','date':'2023.03','kind':'论文v1','label':'角色对话协作','body':'通过角色扮演式对话\n探索多智能体合作','logo':'camel_icon.png','name_text':'CAMEL','paper':'https://arxiv.org/abs/2303.17760','code':'https://github.com/camel-ai/camel','venue':'NeurIPS 2023','in_original':True},
 {'name':'Generative Agents','date':'2023.04','kind':'论文v1','label':'记忆、反思与规划','body':'用记忆、反思与规划\n模拟虚拟角色行为','logo':'generative_agents_icon.png','name_text':'Generative\nAgents','paper':'https://arxiv.org/abs/2304.03442','code':'https://github.com/joonspk-research/generative_agents','venue':'UIST 2023','in_original':True},
 {'name':'MetaGPT','date':'2023.08','kind':'论文v1','label':'流程化多角色协作','body':'以标准流程组织分工\n实现协作式软件开发','logo':'metagpt_icon.png','name_text':'MetaGPT','paper':'https://arxiv.org/abs/2308.00352','code':'https://github.com/FoundationAgents/MetaGPT','venue':'ICLR 2024；本页用论文v1时间，非项目最早代码时间','in_original':True},
 {'name':'Mixture-of-Agents','date':'2024.06','kind':'论文v1','label':'多模型分层聚合','body':'分层聚合多模型回答\n以协作改善生成结果','logo':None,'name_text':'Mixture-of-\nAgents','paper':'https://arxiv.org/abs/2406.04692','code':'https://github.com/togethercomputer/MoA','venue':'本页采用2024年预印本口径','in_original':True},
 {'name':'VideoMind','date':'2025.03','kind':'论文v1','label':'多模态视频推理','body':'结合规划、定位验证\n实现时序定位推理','logo':'videomind_icon.png','name_text':'VideoMind','paper':'https://arxiv.org/abs/2503.13444','code':'https://github.com/yeliudev/VideoMind','venue':'作者技术报告2025-03-17；后续会议时间不作本页年份','in_original':True},
 {'name':'BrowseMaster','date':'2025.08','kind':'论文v1','label':'复杂网页检索','body':'规划与执行模块分工\n结合代码化网页搜索','logo':None,'name_text':'BrowseMaster','paper':'https://arxiv.org/abs/2508.09129','code':'https://github.com/sjtu-sai-agents/Browse-Master','venue':'2025年预印本；不是BrowserMaster或通用GUI点击模型','in_original':True},
 {'name':'Tool-R1','date':'2025.09','kind':'论文v1','label':'强化学习工具使用','body':'用样本高效强化学习\n优化多步工具使用','logo':None,'name_text':'Tool-R1','paper':'https://arxiv.org/abs/2509.12867','code':'https://github.com/YBYBZhang/Tool-R1','venue':'2025年预印本；不是VTool-R1','in_original':True},
]

def rgb(value): return RGBColor.from_string(value)

def no_style(shape):
    st=shape._element.find(qn('p:style'))
    if st is not None: shape._element.remove(st)

def rect(s,x,y,w,h,fill):
    q=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    no_style(q);q.fill.solid();q.fill.fore_color.rgb=rgb(fill);q.line.fill.background()
    return q

def text(s,x,y,w,h,t,size=18,color=INK,bold=False,align=PP_ALIGN.LEFT,white_bg=False):
    q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));q.name='Text: '+t.replace('\n',' ')[:40]
    if white_bg:q.fill.solid();q.fill.fore_color.rgb=rgb('FFFFFF')
    f=q.text_frame;f.clear();f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0
    f.auto_size=MSO_AUTO_SIZE.NONE;f.word_wrap=True
    for i,line in enumerate(t.split('\n')):
        p=f.paragraphs[0] if i==0 else f.add_paragraph();p.text=line;p.alignment=align;p.line_spacing=1.08;p.space_before=Pt(0);p.space_after=Pt(0)
        for r in p.runs:
            r.font.name=FONT;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
            ea=OxmlElement('a:ea');ea.set('typeface',FONT);r._r.get_or_add_rPr().append(ea)
    return q

def line(s,x1,y1,x2,y2,width=1.25,color=RED,arrow=False):
    q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));no_style(q)
    q.line.color.rgb=rgb(color);q.line.width=Pt(width)
    if arrow:
        e=OxmlElement('a:tailEnd');e.set('type','triangle');q._element.spPr.find(qn('a:ln')).append(e)
    return q

def picture(s,path,cx,cy,w,h):
    with Image.open(path) as im:iw,ih=im.size
    scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale
    q=s.shapes.add_picture(str(path),Inches(cx-ww/2),Inches(cy-hh/2),width=Inches(ww),height=Inches(hh))
    q.name='Official asset: '+path.name;return q

def generate():
    p=Presentation();p.slide_width=Inches(W);p.slide_height=Inches(H);s=p.slides.add_slide(p.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('FFFFFF')
    # Reproduce the reference's blue squares, large blue heading, thin rule.
    rect(s,.256,.520,.543,.543,BLUE);rect(s,.524,.811,.362,.362,BLUE)
    text(s,.908,.48,12.05,.67,'LLM Agent 代表性工作演进路线（2022—2025）',31,BLUE,True)
    line(s,.886,1.173,13.063,1.173,.8,INK)
    text(s,.62,1.37,12.10,.39,'从推理与行动出发，扩展到自主执行、记忆反思、协作与复杂任务',19,INK)
    text(s,.63,1.86,12.00,.29,'按首次公开月份排序 · 论文采用预印本v1；AutoGPT采用官方演示月 · 横轴非等比例',12,GRAY)
    axis=4.00;centres=[1.20+i*1.3575 for i in range(9)]
    line(s,.58,axis,12.97,axis,3.0,RED,True)
    layout=[]
    # Upper nodes alternate between two heights; lower descriptions use matching tiers.
    for i,(cx,work) in enumerate(zip(centres,WORKS)):
        high=(i%2==0);name_y=2.81 if high else 3.35;icon_cy=2.50 if high else 2.98
        upper_end=name_y+.41
        if '\n' in work['name_text']:
            name_y-=.10;upper_end=name_y+.70
            if work['logo']:icon_cy-=.08
        line(s,cx,upper_end+.04,cx,axis,1.10)
        label_y=4.65 if high else 5.84
        line(s,cx,axis,cx,label_y-.12,1.10)
        if work['logo']:
            picture(s,ASSETS/work['logo'],cx,icon_cy,1.36 if work['name']=='AutoGPT' else .60,.54)
        # Names remain editable even when there is a genuine logo above them.
        name=text(s,cx-1.06,name_y,2.12,.72 if '\n' in work['name_text'] else .44,work['name_text'],18.5 if '\n' in work['name_text'] else 20.5,BLUE,True,PP_ALIGN.CENTER)
        # Separate endpoint and month label, without an embedded chart or grouped object.
        dot=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(cx-.045),Inches(axis-.045),Inches(.09),Inches(.09));no_style(dot)
        dot.fill.solid();dot.fill.fore_color.rgb=rgb(RED);dot.line.fill.background()
        text(s,cx-.49,4.16,.98,.32,work['date'],14.0,RED,True,PP_ALIGN.CENTER,white_bg=True)
        text(s,cx-1.06,label_y,2.12,.34,work['label'],16.0,RED,True,PP_ALIGN.CENTER,white_bg=True)
        text(s,cx-1.06,label_y+.41,2.12,.95,work['body'],16.0,INK,False,PP_ALIGN.CENTER)
        layout.append({'name':work['name'],'month':work['date'],'node_x_inches':cx,'node_y_inches':axis,'description_tier':'upper' if high else 'lower'})
    # Blue closing strip mirrors the supplied reference, with a restrained academic conclusion.
    text(s,.63,7.00,12.15,.19,'内容据 Farooq et al., Information Fusion 136 (2026), 104444 作者公开时间轴筛选；日期逐项重核。详细来源见备注与 references.md。',8.2,GRAY)
    rect(s,.412,7.16,12.536,.34,BLUE)
    text(s,.60,7.196,12.16,.27,'能力持续拓展，技术路线长期并行；以上为代表性节点，不是线性替代关系。',11.8,'FFFFFF',True,PP_ALIGN.CENTER)
    # Plain, default speaker notes only.
    speaker='建议讲解4—6分钟。先沿红线说明九项工作的首次公开月份，再按能力主线串联，不把横向距离理解为真实时间间隔。\n'
    speaker+='早期ReAct组织推理、行动与观察；AutoGPT展示连续任务执行；CAMEL和Generative Agents分别扩展角色协作及记忆反思；MetaGPT用标准操作流程约束多角色开发；MoA探索多模型分层聚合；VideoMind面向时间定位视频推理；BrowseMaster面向复杂网页信息检索；Tool-R1探索强化学习优化工具使用。\n'
    speaker+='这些方向从早期就交织发展；CAMEL早于Generative Agents，不人为调整顺序以迁就能力叙事。MoA主要是多模型响应聚合，不应等同于具备环境行动的自主Agent。MetaGPT节点采用2023年论文v1，不能说其项目到2024年才出现。BrowseMaster不是BrowserMaster，也不应表述成通用GUI自动化。\n'
    speaker+='原图来源核实：作者官方仓库itsazibfarooq/agenticEvaluation中提供一致内容的timeline.png，README所列DOI与用户提供相同；其图注称这是2023—2025代表性系统的示意时间轴。出版商与TechRxiv正文在本环境返回403，未逐字核定正式版Figure 1图注。所有所选工作均来自原图，没有新增论文。月份依据预印本编号、作者项目和交叉书目；未虚构未直接核对的日级提交日期。\n'
    speaker+='论文来源：https://doi.org/10.1016/j.inffus.2026.104444\n作者时间轴：https://github.com/itsazibfarooq/agenticEvaluation/blob/main/static/images/timeline.png\n'
    for work in WORKS:
        speaker+=f"\n{work['name']}｜{work['date']}｜{work['kind']}｜{work['venue']}\n{work['body'].replace(chr(10),'')}\n{work['paper'] or work['code']}\n官方实现：{work['code']}\n"
    speaker+='\n图标采用官方仓库真实素材；AutoGPT采用当前官方品牌标识作项目识别，不声称它是2023年首发时使用的版本。无专属图标的项目采用可编辑规范文字名称。\n'
    s.notes_slide.notes_text_frame.text=speaker
    p.core_properties.title='LLM Agent 代表性工作演进路线（2022—2025）';p.core_properties.author='左金隆';p.core_properties.subject='代表性工作、首次公开月份与能力演进'
    path=ROOT/'LLM_Agent_Timeline.pptx';p.save(path)
    # Discard optional version-specific metadata inherited from the blank template.
    cleaned=path.with_suffix('.tmp.pptx')
    with ZipFile(path) as zin,ZipFile(cleaned,'w',ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data=zin.read(info.filename)
            if info.filename.endswith('.xml'):
                root=etree.fromstring(data)
                for ext in root.xpath('//*[local-name()="extLst"]'):ext.getparent().remove(ext)
                data=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
            zout.writestr(info,data)
    cleaned.replace(path)
    (ROOT/'timeline_data.json').write_text(json.dumps(WORKS,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'layout_manifest.json').write_text(json.dumps(layout,ensure_ascii=False,indent=2),encoding='utf-8')
    print(path)
    return path

def render(path):
    import fitz
    fontconf=ROOT/'fontconfig.conf'
    env=os.environ.copy()
    if fontconf.exists():env['FONTCONFIG_FILE']=str(fontconf)
    with tempfile.TemporaryDirectory(prefix='llm-agent-render-') as tmp:
        cmd=['soffice','--headless','-env:UserInstallation=file://'+tmp,'--convert-to','pdf','--outdir',str(ROOT),str(path)]
        subprocess.run(cmd,env=env,check=True,timeout=60)
    pdf=fitz.open(ROOT/'LLM_Agent_Timeline.pdf');assert len(pdf)==1
    pdf[0].get_pixmap(matrix=fitz.Matrix(3,3),alpha=False).save(ROOT/'LLM_Agent_Timeline.png')
    print(ROOT/'LLM_Agent_Timeline.png')

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--render',action='store_true');args=ap.parse_args()
    target=generate()
    if args.render:render(target)
